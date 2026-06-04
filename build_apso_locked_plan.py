from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "APSO_Locked_Freeware_Implementation_Plan.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(hdr[i], header, bold=True)
        set_cell_shading(hdr[i], "D9EAF7")
        if widths:
            hdr[i].width = widths[i]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
            if widths:
                cells[i].width = widths[i]
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "EEF6E8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(37, 91, 47)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    r2.font.size = Pt(9)
    doc.add_paragraph()


def setup_styles(doc):
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size, color in [
        ("Title", 22, RGBColor(23, 55, 94)),
        ("Heading 1", 15, RGBColor(23, 55, 94)),
        ("Heading 2", 12, RGBColor(47, 84, 150)),
        ("Heading 3", 10.5, RGBColor(31, 78, 121)),
    ]:
        style = styles[name]
        style.font.name = "Aptos Display" if name in ("Title", "Heading 1") else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = name != "Normal"
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(8 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    setup_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("APSO Locked Freeware Implementation Plan")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Configuration-first, read-only integration baseline for production-grade AI SDLC intelligence")
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(89, 89, 89)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Environment model: dev and prod only | Core development cost: no paid software or paid SaaS dependency")

    add_callout(
        doc,
        "Development Non-Deviation Rule",
        "The APSO team must follow this document as the delivery authority. No alternate framework, paid SaaS dependency, model provider, database, auth platform, object store, workflow engine, observability stack, scanner, or write-back integration may be introduced without formal approval and a versioned update to this document.",
    )

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "APSO will be built as a sellable AI SDLC intelligence platform that extracts evidence from Jira, Bitbucket, AWS CodeCommit, and CI/CD pipelines, then analyzes requirement quality, code quality, coverage, release readiness, and productivity indicators. The original development scope is read-only extraction and analysis. Jira or Bitbucket write-back is not part of the MVP."
    )
    doc.add_paragraph(
        "The stack is locked to free/open-source/self-hosted components for development. Jira, Bitbucket, and AWS are external customer systems and are integrations only; they are not required as paid development dependencies because dev must support fixture/mock mode."
    )

    doc.add_heading("2. Locked Environments", level=1)
    add_table(
        doc,
        ["Environment", "Purpose", "Allowed Integration Mode", "AI Model Policy"],
        [
            ["dev", "Internal development, testing, demos, fixture data, and optional read-only live connector checks.", "Mock/fixture by default; live read-only tokens allowed when available.", "Ollama only: qwen2.5-coder:14b and qwen3-embedding:4b."],
            ["prod", "Customer/self-hosted production deployment.", "Live customer Jira, Bitbucket, AWS CodeCommit, and CI/CD integrations.", "AI Gateway controls provider. Default remains self-hosted/local unless approved."],
        ],
        [Inches(1.0), Inches(2.6), Inches(2.3), Inches(2.5)],
    )

    doc.add_heading("3. Locked Freeware/Open-Source Stack", level=1)
    add_table(
        doc,
        ["Layer", "Locked Tool", "Purpose", "Cost/License Posture"],
        [
            ["Frontend", "Next.js, React, TypeScript, Tailwind CSS, shadcn/ui", "Build APSO web UI, dashboards, findings inbox, settings, reports.", "Free/open source."],
            ["Backend API", "FastAPI, Pydantic, SQLAlchemy, Alembic", "APIs, validation, database access, migrations.", "Free/open source."],
            ["Auth", "Self-hosted Supabase Auth", "User auth, sessions, JWT, tenant identity.", "Self-hosted open-source; no Supabase Cloud dependency."],
            ["Database", "PostgreSQL + pgvector", "System of record, vectors, audit, metrics, normalized SDLC graph.", "Free/open source."],
            ["Object Storage", "MinIO", "Reports, scan artifacts, CI logs, exports, attachments metadata bundles.", "Free/open source AGPLv3; do not modify source without legal review."],
            ["AI Runtime", "Ollama", "Run local models for development and default self-hosted AI.", "Free/open source."],
            ["LLM", "qwen2.5-coder:14b", "Story/code analysis, explanation, recommendations.", "Free model; verify exact license before commercial release."],
            ["Embeddings", "qwen3-embedding:4b", "Semantic search over Jira, PRs, code snippets, logs, docs.", "Free model; verify exact license before commercial release."],
            ["AI Workflow", "LangGraph", "Repeatable multi-step agent workflows.", "Free/open source library."],
            ["RAG", "LlamaIndex", "Retrieval/indexing over tenant-approved data.", "Free/open source library."],
            ["Workflow Engine", "Temporal OSS", "Durable sync, scan, AI analysis, and report workflows.", "Free/open source."],
            ["Cache/Queue", "Valkey", "Cache, locks, rate limits, temporary job state.", "Free/open source BSD."],
            ["Metrics", "Prometheus", "Service and workflow metrics.", "Free/open source."],
            ["Dashboards", "Grafana OSS", "Operational dashboards.", "Free/open source AGPLv3; do not modify source without legal review."],
            ["Logs", "Loki", "Centralized logs.", "Free/open source AGPLv3; do not modify source without legal review."],
            ["Tracing", "Jaeger", "Distributed tracing across services.", "Free/open source."],
            ["Scanners", "Semgrep CE, tree-sitter, Gitleaks, Trivy, OWASP Dependency-Check", "Static analysis, parsing, secret detection, vulnerability scanning.", "Free/open-source/community editions only."],
            ["Deployment", "Docker Compose for dev, k3s for prod", "Repeatable self-hosted deployment.", "Free/open source."],
        ],
        [Inches(1.05), Inches(1.75), Inches(3.3), Inches(2.1)],
    )

    doc.add_heading("4. Configuration-First Rule", level=1)
    doc.add_paragraph(
        "All external systems and frameworks must be configurable. Development must not hard-code Jira site URLs, Bitbucket workspaces, AWS accounts, Supabase endpoints, model names, scanner rules, storage buckets, scoring weights, or environment behavior."
    )
    add_table(
        doc,
        ["Configuration Area", "Must Be Configurable"],
        [
            ["Environment", "dev/prod, URLs, CORS, logging, feature flags."],
            ["Tenant", "tenant name, domain, timezone, modules, retention, AI policy."],
            ["Supabase", "self-hosted URL, JWT secret, anon/service keys, role mapping."],
            ["Postgres", "database URL, schemas, RLS, migrations, backup policy."],
            ["Jira", "base URL/cloud ID, auth token, project keys, field mappings, sync mode."],
            ["Bitbucket", "workspace, repositories, token, branch patterns, scope validation, sync mode."],
            ["AWS", "region, account, CodeCommit repos, CodeBuild/CodePipeline names, read-only role."],
            ["AI Gateway", "provider, model, embedding model, prompt version, output schema, external AI flag."],
            ["Scanners", "rulesets, ignored paths, severity thresholds, enabled tools."],
            ["Storage", "MinIO endpoint, buckets, retention, max file size."],
            ["Observability", "Prometheus scrape, Grafana dashboards, Loki labels, Jaeger sampling."],
        ],
        [Inches(1.7), Inches(6.4)],
    )

    doc.add_heading("5. Original Development Scope", level=1)
    add_bullets(
        doc,
        [
            "Extract and normalize Jira work items, comments, statuses, custom fields, and project metadata.",
            "Extract and normalize Bitbucket repositories, commits, pull requests, branches, pipeline results, logs, tests, and code-insight evidence where readable.",
            "Extract and normalize AWS CodeCommit, CodeBuild, and CodePipeline evidence in read-only mode.",
            "Run Definition Gap, Quality Benchmarking, and Coverage Verification engines.",
            "Generate explainable findings, scores, dashboards, release readiness views, and proof-pack reports.",
            "Keep all external systems read-only for MVP. No Jira or Bitbucket write-back in MVP.",
        ],
    )

    doc.add_heading("6. Jira Integration Scope", level=1)
    doc.add_paragraph(
        "The current Jira API token scopes are sufficient for original read-only development when the token user has access to the target projects."
    )
    add_table(
        doc,
        ["Scope", "Required Now?", "Purpose"],
        [
            ["read:account", "Yes", "Read account/user profile information needed for attribution and connector identity."],
            ["read:jira-work", "Yes", "Read Jira project and issue data, search issues, and read associated work objects permitted to the token user."],
            ["manage:jira-webhook", "No, backlog", "Only needed later if APSO manages Jira webhooks instead of scheduled polling."],
            ["write:jira-work", "No, backlog", "Only needed later if APSO writes comments, updates issues, changes priorities, or creates Jira tickets."],
        ],
        [Inches(1.9), Inches(1.2), Inches(5.0)],
    )
    add_callout(
        doc,
        "Locked Jira MVP Policy",
        "Jira integration is read-only for MVP. APSO must use scheduled polling first. No Jira issue creation, comments, field updates, priority changes, label changes, status transitions, or AI recommendation write-back may be implemented in MVP.",
    )

    doc.add_heading("7. Bitbucket Integration Scope", level=1)
    doc.add_paragraph(
        "For the Bitbucket API token shown in the selection screen, choose only read scopes required for extraction and evidence analysis. Atlassian states that repository read does not include pull requests, so PR read must be selected separately. Pipeline read includes pipeline information such as pipelines, steps, artifacts, logs, tests, and code-insights."
    )
    add_table(
        doc,
        ["Bitbucket Scope", "Select?", "Why APSO Needs It"],
        [
            ["read:account", "Yes", "View user profiles/account context where required by the token flow."],
            ["read:me", "Yes", "Validate the authenticated token owner and connector identity."],
            ["read:user:bitbucket", "Yes", "Read user information for PR authors/reviewers and attribution."],
            ["read:workspace:bitbucket", "Yes", "List/view workspaces the token can access."],
            ["read:project:bitbucket", "Yes", "Read projects and project metadata for grouping repositories."],
            ["read:repository:bitbucket", "Yes", "Read repositories, source code, branches, files, and repository configuration."],
            ["read:pullrequest:bitbucket", "Yes", "Read pull requests, reviews, PR metadata, and PR evidence."],
            ["read:pipeline:bitbucket", "Yes", "Read pipelines, steps, caches, artifacts, logs, tests, and code insights."],
            ["read:test:bitbucket", "Yes", "Read repository/workspace test data where exposed separately in the token UI."],
            ["read:permission:bitbucket", "Optional", "Useful for diagnostics and permission validation; not needed for core extraction if access is already known."],
            ["read:webhook:bitbucket", "Optional/backlog", "Only view existing webhooks. MVP should use polling unless webhook reading is needed for diagnostics."],
            ["read:runner:bitbucket", "No for MVP", "Runner metadata is not required for the first APSO extraction scope."],
            ["read:issue:bitbucket", "No unless Bitbucket Issues are used", "Only needed if the team tracks work in Bitbucket Issues; Jira is the system of record."],
            ["read:wiki:bitbucket", "No unless wiki requirements are used", "Only needed if requirements/docs live in Bitbucket Wiki."],
            ["read:snippet:bitbucket", "No", "Not part of APSO MVP evidence model."],
            ["read:ssh-key:bitbucket", "No", "Not needed for SDLC analysis."],
            ["read:gpg-key:bitbucket", "No", "Not needed for SDLC analysis."],
            ["read:package:bitbucket", "No for MVP", "Only needed if package registry evidence is added later."],
        ],
        [Inches(2.15), Inches(1.2), Inches(4.8)],
    )
    add_callout(
        doc,
        "Locked Bitbucket MVP Policy",
        "Bitbucket integration is read-only for MVP. Do not request write/admin/delete scopes. Do not create branches, commits, pull requests, comments, pipeline runs, pipeline variables, code insights, webhooks, repository settings, or permission changes in MVP.",
    )

    doc.add_heading("8. Backlog: Write-Back And Automation Features", level=1)
    doc.add_paragraph(
        "The following items are explicitly backlog/future-phase capabilities. They explain why write scopes may be useful later, but they are not part of the original development scope."
    )
    add_table(
        doc,
        ["Backlog Item", "Future Permission Needed", "Description"],
        [
            ["Create Jira tickets from AI findings", "write:jira-work", "Allow an approved user to convert an APSO finding into a Jira task/bug."],
            ["Add comments to Jira stories", "write:jira-work", "Write AI recommendations or reviewer notes back to Jira comments after human approval."],
            ["Update Jira priority", "write:jira-work", "Apply approved priority recommendations to Jira."],
            ["Change Jira labels", "write:jira-work", "Apply approved labels such as risk, quality-gap, or needs-nfr."],
            ["Move Jira issue status", "write:jira-work", "Transition issue workflow only after explicit approval."],
            ["Write back AI recommendations into Jira", "write:jira-work", "Persist approved APSO analysis into Jira fields/comments."],
            ["Manage Jira webhooks", "manage:jira-webhook", "Register/refresh/delete Jira webhooks for real-time sync."],
            ["Create Bitbucket PR comments", "Bitbucket write pull request scope", "Comment on pull requests with approved APSO findings."],
            ["Upload Bitbucket code insights", "Bitbucket write pipeline/code-insight scope", "Publish APSO scan results back into Bitbucket."],
            ["Trigger Bitbucket pipelines", "Bitbucket write pipeline scope", "Start/stop/retry pipelines from APSO."],
            ["Manage Bitbucket webhooks", "Bitbucket webhook write/admin scope", "Create/update/delete Bitbucket webhooks for real-time sync."],
        ],
        [Inches(2.35), Inches(1.75), Inches(4.0)],
    )

    doc.add_heading("9. Required Config Files Before Development", level=1)
    add_table(
        doc,
        ["File", "Purpose"],
        [
            ["config/environments.dev.yaml", "All dev URLs, mock/live mode flags, local Ollama, local service endpoints."],
            ["config/environments.prod.yaml", "Production URLs, live integration modes, production secrets references."],
            ["config/integrations.jira.yaml", "Jira project keys, field mappings, polling interval, read-only token reference."],
            ["config/integrations.bitbucket.yaml", "Workspace/repo mapping, branch patterns, required read scopes, polling interval."],
            ["config/integrations.aws.yaml", "AWS region/account, CodeCommit repos, CodeBuild/CodePipeline names, read-only role reference."],
            ["config/ai-gateway.yaml", "Ollama provider, Qwen models, output validation, prompt versions."],
            ["config/scanners.yaml", "Semgrep/tree-sitter/Gitleaks/Trivy/OWASP settings."],
            ["config/scoring-policies.yaml", "Requirement quality, code risk, release readiness scoring weights."],
            ["config/feature-flags.yaml", "Phase-gated features and backlog toggles."],
            ["config/security.yaml", "Credential encryption, JWT, webhook secrets, PII redaction, audit policy."],
        ],
        [Inches(2.6), Inches(5.4)],
    )

    doc.add_heading("10. Phase Plan", level=1)
    add_table(
        doc,
        ["Phase", "Build", "Must Not Include"],
        [
            ["Phase 1: Foundation", "dev/prod config, auth, Postgres/pgvector, MinIO, Valkey, Temporal, AI Gateway, Ollama/Qwen setup.", "Paid model APIs, live write-back, alternate tools."],
            ["Phase 2: Read-Only Integrations", "Jira, Bitbucket, AWS read-only polling, normalized storage, integration health.", "Jira/Bitbucket writes, webhook creation, admin scopes."],
            ["Phase 3: MVP AI Engines", "Definition Gap, Quality Benchmarking, Coverage Verification, AI Findings Inbox.", "Priority automation, contributor scoring, direct provider calls."],
            ["Phase 4: Reporting", "Dashboards, release readiness report, proof pack exports, placeholders for missing metrics.", "Invented productivity metrics."],
            ["Phase 5: Production Hardening", "RLS tests, audit logs, observability, backups, rate limits, security tests.", "Paid SaaS dependencies."],
            ["Phase 6: Backlog/Advanced", "Write-back, webhooks, Priority Engine, Semantic Alignment, Contributor Flow, paid model adapters if approved.", "Anything implemented without scope approval."],
        ],
        [Inches(1.65), Inches(4.0), Inches(2.45)],
    )

    doc.add_heading("11. Acceptance Checklist", level=1)
    add_bullets(
        doc,
        [
            "Only dev and prod environments exist.",
            "All Jira and Bitbucket MVP integrations are read-only.",
            "Jira write-back and Bitbucket write-back are documented only as backlog items.",
            "No paid SaaS dependency is required for development.",
            "All AI calls go through APSO AI Gateway.",
            "No external AI provider is enabled in dev.",
            "All integrations are configuration-driven.",
            "All AI outputs are schema validated before storage.",
            "All business tables include tenant isolation.",
            "All credentials are stored encrypted or referenced through approved secret handling.",
        ],
    )

    doc.add_heading("12. Reference Notes", level=1)
    doc.add_paragraph(
        "Scope guidance is based on Atlassian Jira and Bitbucket Cloud scope documentation checked on 2026-05-31. Bitbucket API token permissions define repository read, pull request read, pipeline read, project read, workspace read, and user read as separate selectable scopes; repository read does not automatically include pull requests. Jira read-only development uses read:account and read:jira-work; write:jira-work and webhook management remain backlog-only."
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("APSO Locked Freeware Implementation Plan | Read-only MVP scope | Configuration-first delivery")

    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
