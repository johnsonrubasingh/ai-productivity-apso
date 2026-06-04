from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TECH_DOC = "docs/project-plan/APSO_Technical_Architecture_and_Tech_Stack_Guide.docx"
BUSINESS_DOC = "docs/project-plan/APSO_Business_Functionality_and_Value_Guide.docx"


BLUE = RGBColor(23, 55, 94)
MID_BLUE = RGBColor(47, 84, 150)
GREEN = RGBColor(37, 91, 47)
GRAY = RGBColor(89, 89, 89)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def margins(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)


def styles(doc):
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size, color in [
        ("Title", 22, BLUE),
        ("Heading 1", 15, BLUE),
        ("Heading 2", 12, MID_BLUE),
        ("Heading 3", 10.5, MID_BLUE),
    ]:
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name in ("Title", "Heading 1") else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)


def title_page(doc, title, subtitle, audience):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run(subtitle)
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = GRAY
    m = doc.add_paragraph()
    m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    m.add_run(f"Audience: {audience}")
    m.add_run("\nLocked environments: dev and prod only")
    m.add_run("\nMVP integration policy: Jira and Bitbucket read-only; AWS mocked until supplied")
    doc.add_paragraph()


def callout(doc, title, body, fill="EEF6E8"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = GREEN
    r.font.size = Pt(10)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.add_run(body).font.size = Pt(9)
    doc.add_paragraph()


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def table(doc, headers, rows, widths=None, header_fill="D9EAF7"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        r.bold = True
        r.font.size = Pt(8.5)
        shade(c, header_fill)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            c.width = widths[i]
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val))
            r.font.size = Pt(8.5)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[i].width = widths[i]
    doc.add_paragraph()
    return t


def component_detail(doc, name, role, apsouse, implementation, configuration, rules):
    doc.add_heading(name, level=2)
    doc.add_paragraph(role)
    table(
        doc,
        ["Aspect", "Guidance"],
        [
            ["Purpose in APSO", apsouse],
            ["Implementation Guidance", implementation],
            ["Configuration", configuration],
            ["Non-Deviation Rules", rules],
        ],
        [Inches(1.55), Inches(6.55)],
        header_fill="EEF6E8",
    )


def footer(doc, text):
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.text = text


def technical_doc():
    doc = Document()
    margins(doc)
    styles(doc)
    title_page(
        doc,
        "APSO Technical Architecture and Tech Stack Guide",
        "Expanded Multi-Page Edition - locked freeware implementation guide for developers and technical leads",
        "Developers, tech leads, architects, DevOps, security reviewers",
    )
    callout(
        doc,
        "Document Authority",
        "This guide is a controlling implementation reference. Development must not introduce alternate frameworks, paid services, direct model calls, write-back integrations, or additional environments without a formal document update.",
    )
    doc.add_page_break()

    doc.add_heading("1. Solution Overview", level=1)
    doc.add_paragraph(
        "APSO is an AI-powered SDLC intelligence platform. It extracts evidence from Jira, Bitbucket, AWS CodeCommit/CodeBuild/CodePipeline, source repositories, CI/CD outputs, and test artifacts. It then analyzes requirement quality, code quality, traceability, release readiness, and productivity indicators using deterministic scanners plus local AI reasoning."
    )
    bullets(
        doc,
        [
            "MVP objective: deliver read-only SDLC intelligence with explainable findings and evidence-backed reports.",
            "Development objective: use only freeware/open-source/self-hosted components and local AI models.",
            "Architecture objective: keep all integrations, models, scanners, scoring policies, and environments configurable.",
            "Commercial objective: build a sellable production-grade product while avoiding paid dependencies during development.",
        ],
    )

    doc.add_heading("2. Locked Environment Model", level=1)
    table(
        doc,
        ["Environment", "Purpose", "Integration Policy", "AI Policy", "Data Policy"],
        [
            ["dev", "Development, internal demos, connector tests, fixture/mocked AWS workflows.", "Jira/Bitbucket read-only; AWS mocked.", "Ollama only with Qwen models.", "Use minimal live data; do not store secrets in Git."],
            ["prod", "Self-hosted production/customer deployment.", "Live read-only integrations unless future write-back is approved.", "AI Gateway controlled; default local/self-hosted.", "Tenant-isolated, audited, backed up."],
        ],
        [Inches(0.8), Inches(1.8), Inches(2.0), Inches(1.7), Inches(1.8)],
    )
    doc.add_page_break()

    doc.add_heading("3. High-Level Architecture", level=1)
    doc.add_paragraph(
        "The product is composed of separate services with clear contracts. Business modules call APSO services, not external providers directly."
    )
    table(
        doc,
        ["Layer", "Components", "Responsibility"],
        [
            ["User Experience", "Next.js web app", "Dashboards, integration setup, findings inbox, reports, admin configuration."],
            ["API Layer", "FastAPI backend", "Tenant-aware APIs, auth validation, RBAC, product workflows, configuration access."],
            ["Integration Layer", "Jira, Bitbucket, AWS connectors", "Read-only extraction, polling, sync checkpoints, rate-limit handling."],
            ["Workflow Layer", "Temporal OSS workers", "Durable jobs for sync, scanning, AI analysis, report generation, retries."],
            ["AI Layer", "APSO AI Gateway, Ollama, Qwen, LangGraph, LlamaIndex", "Model routing, prompts, RAG, structured outputs, AI audit."],
            ["Data Layer", "Postgres, pgvector, MinIO, Valkey", "Relational data, embeddings, artifacts, cache, locks, rate limits."],
            ["Observability", "Prometheus, Grafana, Loki, Jaeger", "Metrics, dashboards, logs, traces, operational diagnostics."],
            ["Security/Governance", "Supabase Auth, RLS, audit logs, secret policies", "Authentication, authorization, tenant isolation, compliance evidence."],
        ],
        [Inches(1.25), Inches(2.25), Inches(4.6)],
    )
    doc.add_page_break()

    doc.add_heading("4. Locked Tech Stack and Purpose", level=1)
    table(
        doc,
        ["Component", "Purpose in APSO", "Why Locked", "Developer Rule"],
        [
            ["Next.js / React / TypeScript", "Build the APSO web application, dashboards, settings, findings inbox, and report screens.", "Strong typed frontend ecosystem; no paid account required.", "Use TypeScript for all application code."],
            ["Tailwind CSS / shadcn/ui", "Implement consistent UI components and styling.", "Fast development with reusable accessible components.", "Do not introduce another UI library without approval."],
            ["FastAPI", "Expose tenant-aware backend APIs.", "High productivity Python API framework, good validation and async support.", "All product APIs live behind FastAPI."],
            ["Pydantic", "Validate request/response payloads and AI structured outputs.", "Critical for safe AI output handling.", "No AI result is stored until schema validation passes."],
            ["SQLAlchemy / Alembic", "ORM/database access and migrations.", "Controlled schema evolution.", "Every schema change requires migration."],
            ["Self-hosted Supabase Auth", "Authentication, JWTs, user/session handling.", "Avoids paid auth service while staying product-ready.", "No alternate auth provider in MVP."],
            ["PostgreSQL", "System of record for tenants, users, integrations, SDLC data, findings, audit.", "Reliable open-source relational database.", "Every business table includes tenant_id."],
            ["pgvector", "Store embeddings for semantic search and alignment.", "Avoids separate vector database cost.", "Use Postgres vector tables for v1."],
            ["MinIO", "Store reports, scan artifacts, logs, coverage files, export bundles.", "Self-hosted S3-compatible storage.", "Store metadata in Postgres, large files in MinIO."],
            ["Ollama", "Run local AI models.", "No paid AI API needed for development.", "All local model calls go through AI Gateway."],
            ["qwen2.5-coder:14b", "Default local LLM for code/story analysis and explanations.", "Strong local coder model for development/testing.", "Do not hard-code model name outside config."],
            ["qwen3-embedding:4b", "Default embedding model for semantic search.", "Local embeddings without paid API.", "Embeddings go through AI Gateway."],
            ["LangGraph", "Orchestrate multi-step AI workflows.", "APSO analysis is step-based, not one prompt.", "Use workflows for repeatable AI engines."],
            ["LlamaIndex", "Build retrieval layer over Jira, PRs, code, logs, docs.", "Simplifies RAG/indexing.", "Retrieve only tenant-approved data."],
            ["Temporal OSS", "Run durable long-running workflows.", "Needed for reliable sync, scan, analysis, report jobs.", "All long jobs must be Temporal workflows."],
            ["Valkey", "Cache, distributed locks, rate limits, temporary state.", "Open-source Redis-compatible option.", "Use for transient data only."],
            ["Prometheus", "Collect service and job metrics.", "Open-source metric standard.", "Every service exposes metrics."],
            ["Grafana OSS", "Operational dashboards.", "Open-source dashboards for prod support.", "Do not modify Grafana source."],
            ["Loki", "Centralized logs.", "Works well with Grafana, self-hosted.", "Never log secrets."],
            ["Jaeger", "Distributed tracing.", "Debug cross-service flows.", "Trace major sync/analysis/report flows."],
            ["Semgrep CE", "Static analysis and security rules.", "Deterministic code findings.", "LLM explains scanner evidence; it does not replace it."],
            ["tree-sitter", "Parse code structure and symbols.", "Language-aware code understanding.", "Use parsing instead of regex for code structure."],
            ["Gitleaks", "Detect committed secrets.", "Essential security scanner.", "Run against repo history/snapshots as configured."],
            ["Trivy", "Scan dependencies, containers, IaC.", "Open-source vulnerability scanner.", "Use in code quality and release readiness."],
            ["OWASP Dependency-Check", "Dependency vulnerability analysis.", "Adds software composition evidence.", "Use scan results as source evidence."],
        ],
        [Inches(1.7), Inches(2.45), Inches(2.0), Inches(1.95)],
    )
    doc.add_page_break()

    doc.add_heading("5. Component Deep Dive", level=1)
    doc.add_paragraph(
        "This section explains how every locked component should be used in APSO. Developers should treat these notes as implementation guidance, not optional background reading."
    )
    component_detail(
        doc,
        "5.1 Next.js, React, and TypeScript",
        "Next.js is the locked web application framework, React is the UI runtime, and TypeScript is the mandatory language for frontend code.",
        "This layer owns APSO's user experience: login flow, tenant/project selection, integration setup, AI Findings Inbox, story quality review, code quality review, coverage traceability, executive dashboards, release readiness views, and report/export screens.",
        "Use route-level organization by product area, typed API clients, reusable UI components, and clear loading/error/empty states. Keep business rules in backend services wherever possible; the frontend should present workflow state and call APIs rather than reimplementing scoring or AI logic.",
        "Read base URLs, feature flags, and environment identity from configuration. The frontend may use the Supabase publishable/anon configuration where appropriate, but it must never receive service role keys, database passwords, integration tokens, or AI provider credentials.",
        "Do not introduce another frontend framework, UI runtime, or untyped JavaScript application code. Do not place Jira, Bitbucket, Supabase service role, or AI model secrets in browser-accessible code.",
    )
    component_detail(
        doc,
        "5.2 Tailwind CSS and shadcn/ui",
        "Tailwind provides utility-first styling and shadcn/ui provides reusable component patterns built on accessible primitives.",
        "These tools give APSO a consistent enterprise product interface without buying a design system. They should be used for tables, buttons, dialogs, tabs, filters, forms, badges, alerts, side navigation, dashboards, and admin settings screens.",
        "Create APSO-specific wrappers for common controls such as severity badges, finding status chips, integration health indicators, score cards, filter bars, and report actions. Keep the UI restrained and operational rather than marketing-heavy.",
        "Theme tokens, colors, spacing, and component variants should be centralized. Feature visibility should come from feature flags, not ad hoc conditionals scattered across components.",
        "Do not introduce Material UI, Ant Design, Bootstrap, Chakra, or another component library in MVP. Avoid one-off styling that bypasses the shared component approach.",
    )
    component_detail(
        doc,
        "5.3 FastAPI Backend",
        "FastAPI is the locked backend API framework for APSO product services.",
        "It owns tenant-aware product APIs, authenticated access, integration configuration, findings retrieval, dashboard data, report generation requests, AI task submission, and administrative operations.",
        "Organize APIs by domain such as auth context, tenants, integrations, sdlc data, findings, reports, and admin. Use dependency injection for tenant context, user context, database session, configuration, and authorization checks.",
        "Backend configuration must load from the approved dev/prod environment files and secret references. API behavior must be feature-flag aware and must not depend on hard-coded tenant/project values.",
        "Do not create sidecar API frameworks or direct database access from the frontend. Do not perform long-running sync/scan/AI jobs inside request threads; submit Temporal workflows instead.",
    )
    component_detail(
        doc,
        "5.4 Pydantic",
        "Pydantic is the locked schema validation layer for API contracts, internal DTOs, and AI outputs.",
        "It protects APSO from malformed API payloads and unreliable model responses. Every AI engine output must be represented by a Pydantic model before it can become an APSO finding, score, report section, or recommendation.",
        "Define explicit schemas for requests, responses, AI task inputs, AI task outputs, connector payloads after normalization, and scoring results. Use strict enums for statuses, severities, evidence types, and lifecycle states.",
        "Validation rules should live near the domain schema. Prompt output schemas must be versioned with the prompt/task definition.",
        "Do not store raw model output as trusted business data. Do not accept flexible untyped JSON for core entities like findings, scores, evidence, or report metrics.",
    )
    component_detail(
        doc,
        "5.5 SQLAlchemy and Alembic",
        "SQLAlchemy is the locked database access layer and Alembic is the locked migration tool.",
        "They manage APSO's relational model: tenants, users, projects, integrations, normalized SDLC artifacts, AI runs, AI findings, evidence, reports, and audit logs.",
        "Use migrations for every schema change. Keep models tenant-aware. Use transactions around sync checkpoints, finding creation, and report snapshots. Prefer explicit relationships and indexes for high-volume tables such as commits, pull requests, test runs, and AI findings.",
        "Database URLs, pool settings, migration mode, and schema names must come from configuration. Migration execution must be controlled and repeatable for dev and prod.",
        "Do not make schema changes manually in production. Do not bypass migrations. Do not create tables without tenant isolation design.",
    )
    component_detail(
        doc,
        "5.6 Self-Hosted Supabase Auth",
        "Supabase Auth is the locked authentication system for user management and session/JWT handling.",
        "It provides sign-in, user identity, password/session lifecycle, and token issuance. APSO's backend validates tokens and maps users to tenants and roles before allowing access to data or actions.",
        "Use Supabase Auth for identity, but keep APSO authorization decisions in APSO's own tenant/role tables. The frontend can use public auth configuration, while privileged admin operations must happen only through backend services.",
        "Auth URL, anon/publishable key, service role key reference, JWT secret reference, redirect URLs, and role mapping must be configured. Service role credentials must be backend-only.",
        "Do not expose service role keys to the frontend. Do not rely on frontend-only authorization. Do not introduce another auth provider in MVP.",
    )
    component_detail(
        doc,
        "5.7 PostgreSQL",
        "PostgreSQL is the locked system-of-record database.",
        "It stores durable APSO data including tenants, users, roles, project mappings, integration metadata, normalized Jira/Bitbucket/AWS artifacts, findings, reports, scoring results, and audit history.",
        "Design tables around normalized evidence and traceability. Large files belong in MinIO, not Postgres. Use indexes for tenant_id, project_id, external IDs, timestamps, status fields, and lookup relationships.",
        "Connection strings, pool settings, backup policy, schema names, and migration user must be configured. Production must have backup and restore procedures.",
        "Do not use a second application database in MVP. Do not store raw secrets in database tables unless encrypted through approved secret handling.",
    )
    component_detail(
        doc,
        "5.8 pgvector",
        "pgvector is the locked vector storage extension inside PostgreSQL.",
        "It stores embeddings for Jira stories, requirements, PR comments, code snippets, test summaries, release notes, and documents so APSO can perform semantic retrieval and alignment checks without a separate vector database.",
        "Create vector tables with tenant_id, source artifact reference, chunk metadata, embedding model version, and timestamps. Retrieval must always filter by tenant and project before similarity ranking.",
        "Embedding dimension and model name must match the configured embedding model. Re-embedding strategy must be documented when the embedding model changes.",
        "Do not add Pinecone, Weaviate, Milvus, or another vector database in MVP. Do not run cross-tenant vector searches.",
    )
    doc.add_page_break()
    component_detail(
        doc,
        "5.9 MinIO",
        "MinIO is the locked self-hosted object storage service.",
        "It stores large and generated artifacts: scan outputs, coverage files, CI/CD logs, report exports, proof packs, temporary analysis bundles, and attachment-derived processing artifacts.",
        "Store object metadata and tenant ownership in Postgres. Use deterministic bucket naming, tenant/project prefixes, retention policies, and signed/internal URLs through backend services.",
        "Endpoint, bucket names, access key reference, secret key reference, max file size, retention, and allowed file types must be configured.",
        "Do not store large artifacts in Postgres. Do not expose MinIO credentials to the frontend. Do not modify MinIO source code without legal review because of AGPL obligations.",
    )
    component_detail(
        doc,
        "5.10 Ollama",
        "Ollama is the locked local AI runtime for development and default self-hosted inference.",
        "It lets APSO run local LLM and embedding models without paid API usage. In dev, every AI engine uses Ollama through the APSO AI Gateway.",
        "The AI Gateway should call Ollama for chat/generation and embeddings. Timeout, retry policy, model names, base URL, and task routing must be configured.",
        "Configure Ollama base URL, chat model, embedding model, concurrency limits, timeout, and maximum input size. Keep external providers disabled in dev.",
        "Do not call Ollama directly from product modules. Do not hard-code model names. Do not assume local model output is reliable without validation.",
    )
    component_detail(
        doc,
        "5.11 qwen2.5-coder:14b",
        "qwen2.5-coder:14b is the locked local LLM for APSO development.",
        "It is used for story gap explanations, code quality explanations, semantic reasoning, release risk summaries, and report narrative generation during development and local demos.",
        "Use deterministic scanner/retrieval evidence first, then ask the model to reason over that evidence. Keep prompts task-specific and output schemas strict. Prefer smaller bounded tasks over large vague prompts.",
        "The model name, temperature, max tokens, context limit, retry behavior, and prompt version must be configured per AI task.",
        "Do not use the model as the source of truth for metrics. Do not let it invent defects, hours saved, productivity gains, or facts that are not supported by evidence.",
    )
    component_detail(
        doc,
        "5.12 qwen3-embedding:4b",
        "qwen3-embedding:4b is the locked local embedding model.",
        "It converts text/code artifacts into vectors for semantic search, similarity matching, duplicate detection, and alignment analysis.",
        "Embed curated chunks, not uncontrolled full repositories. Track source artifact, chunk boundaries, tenant, project, embedding model version, and source hash so embeddings can be refreshed when inputs change.",
        "Configure embedding model name, dimension, chunk size, overlap, allowed artifact types, and re-indexing policy.",
        "Do not mix embeddings from different models in the same index without model metadata. Do not embed secrets or excluded files.",
    )
    component_detail(
        doc,
        "5.13 LangGraph",
        "LangGraph is the locked framework for orchestrating multi-step AI workflows.",
        "APSO analysis requires repeatable steps: collect evidence, retrieve context, run scanners, call AI task, validate output, create findings, and update workflow state. LangGraph models these steps explicitly.",
        "Use LangGraph for AI engine workflows such as Definition Gap, Quality Benchmarking, Coverage Verification, and Release Readiness. Each node should have a clear input/output contract.",
        "Workflow/task names, prompts, model routing, retries, and feature flags must be configurable. Persist run IDs and correlate them with AI Gateway audits.",
        "Do not implement important AI engines as one-off scripts or untracked prompt chains. Do not bypass validation nodes.",
    )
    component_detail(
        doc,
        "5.14 LlamaIndex",
        "LlamaIndex is the locked retrieval and indexing layer.",
        "It helps APSO find relevant Jira stories, prior defects, PR comments, code snippets, logs, coverage summaries, and documentation before the AI Gateway generates findings.",
        "Use LlamaIndex to build tenant-scoped indexes over approved data. Retrieval should return source references so every AI answer can cite evidence.",
        "Configure data sources, chunking rules, embedding provider, index refresh policy, tenant filters, and top-k retrieval limits.",
        "Do not retrieve cross-tenant data. Do not pass arbitrary unfiltered repository content to the model. Do not use retrieval results without source references.",
    )
    component_detail(
        doc,
        "5.15 Temporal OSS",
        "Temporal OSS is the locked durable workflow engine.",
        "It runs long jobs reliably: Jira sync, Bitbucket sync, AWS mock/live sync, repository analysis, scanner execution, AI analysis, embedding refresh, and report generation.",
        "Use Temporal workflows for operations that can fail, retry, resume, or take longer than an API request. Activities should be idempotent and checkpointed.",
        "Configure namespace, task queues, retry policies, timeouts, worker concurrency, and dead-letter handling.",
        "Do not run long sync/scan/AI operations in API request handlers. Do not create a second workflow engine in MVP.",
    )
    component_detail(
        doc,
        "5.16 Valkey",
        "Valkey is the locked Redis-compatible cache and coordination service.",
        "It supports short-lived caching, distributed locks, rate limiting, job coordination, and temporary state for API and worker operations.",
        "Use Valkey for transient data only. Common uses include connector rate-limit state, lock keys for repository scans, cached dashboard fragments, and webhook/polling deduplication.",
        "Configure endpoint, password/TLS, key prefixes, TTLs, lock timeouts, and rate-limit windows.",
        "Do not treat Valkey as durable storage. Do not store secrets or tenant data that must survive restarts only in Valkey.",
    )
    doc.add_page_break()
    component_detail(
        doc,
        "5.17 Prometheus",
        "Prometheus is the locked metrics collector.",
        "It collects operational and product-health metrics from APSO services, workers, AI Gateway, connectors, scanners, and infrastructure.",
        "Expose metrics for API latency, sync duration, AI run duration, validation failures, scanner failures, queue depth, database pool usage, and integration error rates.",
        "Configure scrape targets, labels, retention, alert thresholds, and service names consistently across dev and prod.",
        "Do not rely only on logs for production diagnosis. Do not expose sensitive labels or tokens in metrics.",
    )
    component_detail(
        doc,
        "5.18 Grafana OSS",
        "Grafana OSS is the locked dashboarding tool for observability.",
        "It gives operators and tech leads visual dashboards for API health, worker health, integration sync, AI Gateway performance, scanner stability, and infrastructure usage.",
        "Create dashboards as versioned JSON/provisioned config where possible. Include panels for error rate, latency, queue backlog, sync failures, AI validation failures, and storage/database health.",
        "Configure admin credentials, data sources, dashboards, folders, and access roles. Keep Grafana credentials outside Git.",
        "Do not modify Grafana source. Do not use Grafana as the customer-facing product UI; it is operational tooling.",
    )
    component_detail(
        doc,
        "5.19 Loki",
        "Loki is the locked centralized log store.",
        "It collects logs from APIs, workers, AI Gateway, integrations, scanners, and supporting services so failures can be investigated across the system.",
        "Use structured logs with tenant-safe correlation IDs, workflow IDs, integration IDs, and request IDs. Mask secrets before logging.",
        "Configure log labels, retention, ingestion endpoints, and service names. Keep logs useful but not overly verbose in prod.",
        "Do not log API tokens, service role keys, database URLs with passwords, JWTs, or raw secret-bearing payloads.",
    )
    component_detail(
        doc,
        "5.20 Jaeger",
        "Jaeger is the locked distributed tracing tool.",
        "It shows the full path of a request or workflow across frontend, API, Temporal, connectors, scanners, AI Gateway, database, and object storage.",
        "Instrument important paths such as project sync, repository scan, story analysis, release readiness computation, and report generation.",
        "Configure OpenTelemetry exporters, trace sampling, service names, and correlation IDs.",
        "Do not include secrets or full sensitive payloads in trace attributes. Do not skip tracing for long-running workflows.",
    )
    component_detail(
        doc,
        "5.21 Semgrep CE",
        "Semgrep Community Edition is the locked static analysis scanner.",
        "It detects code patterns related to security, correctness, framework misuse, and maintainability. APSO uses it as deterministic source evidence for quality findings.",
        "Run Semgrep against configured repositories/branches and store normalized findings with file, line, rule, severity, and message. AI may summarize and prioritize the findings but cannot replace scanner output.",
        "Configure rulesets, custom rules, severity thresholds, ignored paths, languages, and scan timeout.",
        "Do not treat LLM opinions as scanner findings. Do not require Semgrep Cloud in MVP.",
    )
    component_detail(
        doc,
        "5.22 tree-sitter",
        "tree-sitter is the locked code parsing library.",
        "It gives APSO language-aware structure such as functions, classes, imports, symbols, changed methods, and complexity candidates.",
        "Use tree-sitter to extract symbol metadata, changed code regions, dependency relationships, and code chunks for embeddings/RAG.",
        "Configure supported languages, parser versions, file include/exclude patterns, and maximum file size.",
        "Do not parse code structure with fragile regex when tree-sitter supports the language. Do not index generated/vendor files.",
    )
    component_detail(
        doc,
        "5.23 Gitleaks",
        "Gitleaks is the locked secret detection scanner.",
        "It detects committed secrets such as API keys, credentials, tokens, private keys, and sensitive connection strings.",
        "Run it against configured repository snapshots/history based on policy. Store findings as security evidence with masked secret value, rule ID, file path, and commit/reference.",
        "Configure allowlists, ignored paths, scan depth, redaction, and severity mapping.",
        "Do not display full secrets in UI, logs, exports, or AI prompts. Do not send secret-like findings to external models.",
    )
    component_detail(
        doc,
        "5.24 Trivy",
        "Trivy is the locked vulnerability and configuration scanner.",
        "It scans dependencies, containers, lockfiles, and infrastructure-as-code files for known vulnerabilities and misconfigurations.",
        "Run Trivy as part of repository quality analysis and release readiness checks. Normalize CVEs, package names, severities, fixed versions, and artifact references.",
        "Configure scan targets, vulnerability database cache, severity thresholds, ignored vulnerabilities, and timeout.",
        "Do not fail APSO workflows solely because Trivy finds an issue; convert results into findings and release risk evidence based on policy.",
    )
    component_detail(
        doc,
        "5.25 OWASP Dependency-Check",
        "OWASP Dependency-Check is the locked software composition analysis tool.",
        "It provides an additional dependency vulnerability view, especially useful for Java/JVM and common package ecosystems.",
        "Run it where dependency manifests are present and store normalized dependency evidence. Use it together with Trivy rather than replacing one with the other.",
        "Configure data cache, ecosystem support, suppression files, severity thresholds, and scan locations.",
        "Do not introduce paid SCA tooling in MVP. Do not expose raw dependency reports without tenant controls.",
    )
    component_detail(
        doc,
        "5.26 Docker Compose and k3s",
        "Docker Compose is locked for development orchestration and k3s is locked for production self-hosted Kubernetes deployment.",
        "Compose makes dev setup repeatable for APSO services and dependencies. k3s gives a lightweight production path without managed Kubernetes cost.",
        "Create service definitions for frontend, backend, AI Gateway, workers, Postgres/Supabase components as applicable, MinIO, Valkey, Temporal, Prometheus, Grafana, Loki, Jaeger, and Ollama integration.",
        "Configuration must come from dev/prod files and environment variables. Production manifests should reference secret stores rather than plain values.",
        "Do not introduce managed Kubernetes, Docker Swarm, or another deployment platform in MVP unless formally approved.",
    )
    component_detail(
        doc,
        "5.27 Jira Connector",
        "The Jira connector is the locked source for work item and requirement evidence.",
        "It extracts projects, epics, stories, tasks, bugs, statuses, comments, labels, custom fields, sprint/story-point fields where accessible, and status history for APSO analysis.",
        "Implement scheduled read-only polling first. Normalize external IDs, timestamps, issue hierarchy, comments, status transitions, and field mappings into APSO tables.",
        "Configure base URL, username, token reference, project keys, issue types, custom field mappings, polling interval, and sync start date.",
        "Do not request or use Jira write scopes in MVP. Do not create issues, comments, labels, priority changes, or transitions.",
    )
    component_detail(
        doc,
        "5.28 Bitbucket Connector",
        "The Bitbucket connector is the locked source for repository, pull request, and pipeline evidence.",
        "It extracts workspaces, projects, repositories, branches, commits, pull requests, reviewers, pipeline runs, logs, artifacts, tests, and code insight evidence where readable.",
        "Implement read-only polling for the configured web/API repositories and develop branch first. Normalize commits, PRs, reviews, branch names, build statuses, and test evidence.",
        "Configure workspace, username, token reference, repositories, branches, include/exclude paths, polling interval, and required read scopes.",
        "Do not request or use write/admin/delete scopes in MVP. Do not create PR comments, trigger pipelines, manage webhooks, or change repository settings.",
    )
    component_detail(
        doc,
        "5.29 AWS Mock and Future AWS Connector",
        "AWS integration is deferred until the AWS user/role is supplied, so development must use a mock connector.",
        "The mock connector allows APSO to build the CodeCommit/CodeBuild/CodePipeline data model, UI, and workflows without paid AWS usage or live AWS credentials.",
        "Define the same internal normalized contract that the future live AWS connector will use: repositories, branches, commits, build runs, pipeline executions, artifacts, statuses, and timestamps.",
        "Configure AWS mode as mock in dev and mock_until_credentials_supplied in prod until approved credentials exist.",
        "Do not create AWS resources, IAM users, deployments, or write operations from APSO during MVP.",
    )
    doc.add_page_break()

    doc.add_heading("6. AI Gateway Architecture", level=1)
    doc.add_paragraph(
        "The AI Gateway is mandatory. It is the only component allowed to communicate with Ollama or any future model provider."
    )
    table(
        doc,
        ["Contract", "Description"],
        [
            ["Input schema", "Defines exactly what data a task accepts, including tenant, project, artifact references, and evidence."],
            ["Output schema", "Defines JSON result structure for findings, scores, summaries, confidence, evidence, and next actions."],
            ["Provider adapter", "Contains provider-specific implementation for Ollama now and optional paid providers later."],
            ["Prompt registry", "Stores prompt versions outside business code."],
            ["Validation", "Pydantic validation runs before any model output is stored."],
            ["Audit", "Every run records task, model, prompt version, latency, status, evidence references, and validation result."],
            ["Privacy", "External providers are disabled in MVP; future external calls require tenant approval and redaction."],
        ],
        [Inches(1.7), Inches(6.4)],
    )
    bullets(
        doc,
        [
            "Business modules call task contracts such as requirement_gap_analysis or release_risk_summary.",
            "Business modules never call provider SDKs directly.",
            "Model names live in configuration, not code.",
            "Provider switching is done inside the gateway only.",
        ],
    )
    doc.add_page_break()

    doc.add_heading("7. Data Architecture", level=1)
    table(
        doc,
        ["Schema", "Primary Data", "Notes"],
        [
            ["core", "tenants, users, roles, projects, feature flags", "Tenant identity and product configuration."],
            ["integrations", "connectors, credentials references, sync jobs, checkpoints", "Never store plain-text secrets."],
            ["sdlc", "work items, repositories, commits, PRs, builds, deployments, tests, coverage", "Normalized Jira/Bitbucket/AWS data."],
            ["ai", "ai runs, findings, evidence, scores, prompt versions, feedback", "Every finding must be explainable."],
            ["reports", "proof packs, exports, report snapshots", "Large files stored in MinIO."],
            ["audit", "audit events, access events, config changes", "Used for governance and support."],
        ],
        [Inches(1.2), Inches(3.3), Inches(3.6)],
    )
    callout(
        doc,
        "Tenant Isolation Rule",
        "Every business table must include tenant_id. API access, workers, exports, AI runs, and object storage references must enforce tenant isolation.",
        "FFF2CC",
    )
    doc.add_page_break()

    doc.add_heading("8. Integration Architecture", level=1)
    table(
        doc,
        ["System", "MVP Mode", "Extracted Data", "Not Allowed in MVP"],
        [
            ["Jira", "Live read-only polling", "Projects, issues, issue types, comments, statuses, custom fields, sprints where accessible.", "Creating issues, comments, labels, priority changes, status transitions."],
            ["Bitbucket", "Live read-only polling", "Workspaces, projects, repos, branches, commits, PRs, reviewers, pipelines, logs, tests, artifacts where accessible.", "Creating PR comments, triggering pipelines, modifying repository settings."],
            ["AWS", "Mock until user supplied", "CodeCommit, CodeBuild, CodePipeline evidence once credentials exist.", "AWS write operations, deployments, IAM changes."],
            ["GitHub", "Code repository only", "APSO source control and CI if configured.", "Not a product data source unless later approved."],
        ],
        [Inches(1.0), Inches(1.5), Inches(3.6), Inches(2.0)],
    )
    doc.add_page_break()

    doc.add_heading("9. Core AI Engines", level=1)
    table(
        doc,
        ["Engine", "Inputs", "Processing", "Outputs"],
        [
            ["Definition Gap Engine", "Jira stories, acceptance criteria, comments, custom fields.", "Rule checks plus AI review for ambiguity, missing persona, missing NFR, weak testability.", "Requirement quality score and gap findings."],
            ["Quality Benchmarking Engine", "Repository snapshot, PRs, commits, scanner outputs.", "Semgrep/tree-sitter/Gitleaks/Trivy/Dependency-Check plus AI explanation.", "Code risk score and quality findings."],
            ["Coverage Verification Engine", "Jira work items, commits, PRs, tests, coverage, CI/CD evidence.", "Traceability matching and coverage confidence analysis.", "Coverage gaps and release readiness evidence."],
            ["Release Readiness", "Open findings, build/test results, story status, PR status.", "Aggregate policy checks and risk summarization.", "Release readiness score and executive summary."],
        ],
        [Inches(1.55), Inches(2.2), Inches(2.8), Inches(1.55)],
    )
    doc.add_page_break()

    doc.add_heading("10. Workflow Design", level=1)
    numbered(
        doc,
        [
            "Tenant admin configures Jira, Bitbucket, Supabase, and project mapping.",
            "Temporal starts scheduled read-only sync workflows.",
            "Connectors extract source data and normalize it into Postgres.",
            "Scanner workflows collect deterministic code/security/test evidence.",
            "LlamaIndex retrieves relevant tenant-approved context.",
            "LangGraph executes AI engine workflows through the AI Gateway.",
            "Pydantic validates structured outputs.",
            "Findings, evidence, and scores are stored with audit metadata.",
            "Dashboards and reports present actionable insights to users.",
        ],
    )

    doc.add_heading("11. Security and Governance", level=1)
    bullets(
        doc,
        [
            "Secrets must never be committed; use ignored env files or encrypted local secret storage.",
            "Service role keys are backend-only and must never be exposed to the frontend.",
            "No external AI provider is enabled in development.",
            "Scanner results are evidence; LLMs may explain and prioritize but not invent facts.",
            "All AI findings must include source references.",
            "All write-back features remain backlog until separately approved.",
            "Logs must mask tokens, database passwords, JWTs, and connection strings.",
            "Object storage keys must be tenant scoped through metadata and access controls.",
        ],
    )

    doc.add_heading("12. Development Phases", level=1)
    table(
        doc,
        ["Phase", "Build", "Exit Criteria"],
        [
            ["1 Foundation", "Repo scaffold, config, auth, Postgres/pgvector, MinIO, Valkey, Temporal, AI Gateway, Ollama.", "User logs in; AI Gateway calls Qwen; sample workflow runs."],
            ["2 Read-Only Integrations", "Jira and Bitbucket connectors; AWS mock connector; sync dashboards.", "Data extracted and normalized without writes."],
            ["3 MVP AI Engines", "Definition Gap, Quality Benchmarking, Coverage Verification, Findings Inbox.", "Explainable findings generated from real evidence."],
            ["4 Reporting", "Dashboards, release readiness, proof pack export.", "Business-friendly reports generated with no fabricated metrics."],
            ["5 Production Hardening", "RLS tests, audits, backup, observability, rate limits, security tests.", "Ready for controlled pilot."],
            ["6 Backlog", "Write-back, webhooks, advanced engines, paid model adapters if approved.", "Only after formal scope approval."],
        ],
        [Inches(1.2), Inches(3.7), Inches(3.2)],
    )

    doc.add_heading("13. Developer Acceptance Rules", level=1)
    bullets(
        doc,
        [
            "No component may bypass the AI Gateway.",
            "No table may store tenant data without tenant_id.",
            "No connector may require write permission in MVP.",
            "No score may be generated without evidence or a documented placeholder state.",
            "No secret may appear in source control or logs.",
            "No paid dependency may be introduced into development.",
            "No additional environment may be added.",
            "No backlog feature may be pulled into MVP without approval.",
        ],
    )

    footer(doc, "APSO Technical Architecture and Tech Stack Guide | Locked freeware stack | Read-only MVP")
    doc.save(TECH_DOC)


def business_doc():
    doc = Document()
    margins(doc)
    styles(doc)
    title_page(
        doc,
        "APSO Business Functionality and Value Guide",
        "Stakeholder guide to the AI SDLC intelligence application",
        "Business sponsors, delivery leaders, product owners, program managers, customer stakeholders",
    )
    callout(
        doc,
        "Business Purpose",
        "APSO is designed to prove how AI can improve SDLC productivity, quality, release confidence, and process efficiency using evidence from actual delivery systems.",
    )

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "APSO turns delivery data into actionable intelligence. Instead of showing only sprint metrics or static dashboards, APSO identifies requirement gaps, code risks, coverage gaps, and release readiness issues before they become rework, defects, or delivery delays."
    )
    bullets(
        doc,
        [
            "For leadership: visibility into quality, productivity, release risk, and improvement evidence.",
            "For product owners: better stories, clearer acceptance criteria, and fewer ambiguous requirements.",
            "For engineering leads: earlier detection of code risk, test gaps, security issues, and review bottlenecks.",
            "For customers: a stronger demo and product narrative showing measurable AI impact.",
        ],
    )

    doc.add_heading("2. Business Problem", level=1)
    doc.add_paragraph(
        "The original dashboard demo showed SDLC reporting, but stakeholders expected AI capabilities that actively improve delivery. APSO addresses that expectation by connecting requirements, code, tests, pipelines, and releases into one intelligence layer."
    )
    table(
        doc,
        ["Current Pain", "Business Impact", "APSO Response"],
        [
            ["Ambiguous requirements", "Rework, missed expectations, delayed acceptance.", "Definition Gap Engine finds missing acceptance criteria, unclear terms, missing personas, and missing NFRs."],
            ["Late quality discovery", "Defects found after development or testing.", "Quality Benchmarking Engine catches risks earlier using scanners and AI explanation."],
            ["Weak traceability", "Hard to prove whether a story has code, tests, and CI evidence.", "Coverage Verification links Jira items to commits, PRs, test results, and pipelines."],
            ["Dashboard-only reporting", "Leadership sees status but not actionable improvement.", "Findings inbox and proof packs show risks, actions, and measurable outcomes."],
            ["Unproven AI productivity", "Stakeholders do not see evidence of impact.", "Evidence framework tracks hours saved, rework avoided, defects prevented, and lead-time improvements."],
        ],
        [Inches(2.0), Inches(2.2), Inches(3.9)],
    )

    doc.add_heading("3. Product Vision", level=1)
    doc.add_paragraph(
        "APSO should be positioned as an AI SDLC Control Tower: a product that continuously reviews delivery artifacts, identifies risks, recommends actions, and creates evidence that AI is improving delivery quality and productivity."
    )
    callout(
        doc,
        "Positioning Statement",
        "APSO helps delivery organizations reduce rework, improve release confidence, and prove AI-driven productivity gains by analyzing requirements, code, tests, and pipelines with explainable evidence.",
        "D9EAF7",
    )

    doc.add_heading("4. Main User Groups", level=1)
    table(
        doc,
        ["User Group", "What They Need", "APSO Value"],
        [
            ["Executive Sponsor", "Confidence that AI improves delivery outcomes.", "Proof packs, trends, productivity and quality indicators."],
            ["Delivery Manager", "Visibility into risks, blockers, readiness, and team flow.", "Project health, release readiness, findings by severity."],
            ["Product Owner", "Better stories and fewer missed requirements.", "Story quality scoring and requirement gap recommendations."],
            ["Tech Lead", "Earlier code quality and test coverage visibility.", "Code risk findings, scanner evidence, coverage traceability."],
            ["Developer", "Actionable feedback, not vague dashboards.", "Evidence-based findings with source links and suggested fixes."],
            ["QA Lead", "Clear traceability between stories and tests.", "Coverage gaps, test evidence, release readiness checks."],
            ["Auditor/Reviewer", "Traceability and governance evidence.", "Audit logs, evidence references, report snapshots."],
        ],
        [Inches(1.5), Inches(2.8), Inches(3.8)],
    )

    doc.add_heading("5. Core Application Functionality", level=1)
    table(
        doc,
        ["Functionality", "What It Does", "Business Outcome"],
        [
            ["Integration Setup", "Connects Jira, Bitbucket, and later AWS in read-only mode.", "Fast onboarding without changing customer systems."],
            ["SDLC Data Hub", "Normalizes stories, commits, PRs, pipelines, tests, coverage, and findings.", "One evidence base across delivery lifecycle."],
            ["AI Findings Inbox", "Central queue for gaps, risks, and recommendations.", "Teams can act on findings instead of reading passive reports."],
            ["Requirement Quality Review", "Reviews stories for completeness, ambiguity, testability, and NFR coverage.", "Less rework and clearer development scope."],
            ["Code Quality Review", "Combines static scanners with AI explanations.", "Earlier detection of defects, smells, vulnerabilities, and maintainability risks."],
            ["Coverage Verification", "Checks whether stories are backed by code, PRs, tests, and CI evidence.", "Better release confidence and traceability."],
            ["Release Readiness", "Aggregates open risks, pipeline status, tests, coverage, and unresolved findings.", "Clear go/no-go support for releases."],
            ["Proof Pack Generator", "Creates stakeholder reports with evidence and placeholders for missing actual metrics.", "Supports demos, steering meetings, and customer value proof."],
            ["Audit and Governance", "Tracks user actions, AI runs, configuration changes, and report generation.", "Enterprise trust and supportability."],
        ],
        [Inches(1.8), Inches(3.2), Inches(3.1)],
    )

    doc.add_heading("6. MVP Functionality", level=1)
    doc.add_paragraph(
        "The MVP should be sellable but focused. It should prove the strongest value first: quality, gaps, coverage, and evidence."
    )
    bullets(
        doc,
        [
            "Read-only Jira ingestion for work items, comments, status, custom fields, and requirement text.",
            "Read-only Bitbucket ingestion for repositories, commits, pull requests, pipelines, tests, and evidence.",
            "Mocked AWS connector until AWS user/role is provided.",
            "Definition Gap Engine for story quality.",
            "Quality Benchmarking Engine for code/security risk.",
            "Coverage Verification Engine for traceability and release confidence.",
            "AI Findings Inbox with lifecycle states.",
            "Executive dashboard, project health dashboard, and release readiness report.",
            "Proof pack export for business stakeholders.",
        ],
    )

    doc.add_heading("7. Out of Scope for MVP", level=1)
    table(
        doc,
        ["Backlog Item", "Why Deferred"],
        [
            ["Jira write-back", "MVP should be safe read-only extraction. Write-back requires extra approval and permissions."],
            ["Bitbucket write-back", "PR comments, code insights, and pipeline triggers should wait until findings quality is proven."],
            ["Advanced contributor analytics", "Requires careful governance to avoid surveillance concerns."],
            ["Paid model providers", "Development must avoid paid AI usage."],
            ["Live AWS integration", "AWS user will be supplied later; dev uses mocks first."],
            ["Billing/subscription module", "Not needed to prove product value in MVP."],
            ["Advanced benchmarking", "Requires larger historical data volume and calibrated metrics."],
        ],
        [Inches(2.5), Inches(5.6)],
    )

    doc.add_heading("8. How APSO Works: Business Flow", level=1)
    numbered(
        doc,
        [
            "A tenant admin connects Jira and Bitbucket with read-only credentials.",
            "APSO synchronizes project, story, repository, pull request, pipeline, and test evidence.",
            "APSO analyzes stories for gaps and ambiguity.",
            "APSO analyzes code and pipeline evidence for quality and release risk.",
            "APSO links requirements to code, PRs, tests, and build evidence.",
            "APSO creates explainable findings with severity, confidence, and source links.",
            "Teams review findings, mark actions, and resolve risks.",
            "Leaders view dashboards and generate proof packs showing delivery improvement evidence.",
        ],
    )

    doc.add_heading("9. AI Engines Explained for Stakeholders", level=1)
    table(
        doc,
        ["AI Engine", "Plain-English Explanation", "Example Finding"],
        [
            ["Definition Gap Engine", "Acts like a requirements reviewer that checks whether a story is clear enough to build and test.", "Story lacks acceptance criteria for failure scenarios and has undefined user persona."],
            ["Quality Benchmarking Engine", "Acts like a code quality reviewer that combines scanner evidence with AI explanation.", "PR introduces high-complexity method and missing validation around external input."],
            ["Coverage Verification Engine", "Checks whether delivery work has evidence across code, tests, and pipeline results.", "Story marked done but no linked test evidence or successful pipeline run found."],
            ["Release Readiness Engine", "Summarizes whether the release is safe to proceed based on unresolved risks and evidence.", "Release has 3 high-risk open findings and 2 failed pipeline runs."],
            ["Future Priority Engine", "Will recommend backlog priority based on value, dependency, effort, and risk.", "Deferred to later phase."],
            ["Future Semantic Alignment Engine", "Will compare requirements, code, tests, and release notes for drift.", "Deferred to later phase."],
        ],
        [Inches(1.9), Inches(3.35), Inches(2.85)],
    )

    doc.add_heading("10. Key Dashboards and Screens", level=1)
    table(
        doc,
        ["Screen", "Purpose", "Primary Audience"],
        [
            ["Executive Dashboard", "Shows delivery health, quality trends, release readiness, proof metrics.", "Executives, sponsors."],
            ["Project Health Dashboard", "Shows project-level findings, risks, sync health, readiness.", "Delivery managers, tech leads."],
            ["Findings Inbox", "Shows actionable AI findings with lifecycle management.", "Product, engineering, QA."],
            ["Story Quality View", "Shows requirement gaps, ambiguity, missing NFRs, testability issues.", "Product owners, BAs, scrum leads."],
            ["Code Quality View", "Shows scanner-backed risks and AI explanations.", "Tech leads, developers."],
            ["Coverage Traceability View", "Shows story-to-code-to-test-to-build evidence.", "QA leads, release managers."],
            ["Proof Pack Export", "Creates stakeholder-ready report.", "Business sponsors, customer stakeholders."],
        ],
        [Inches(2.0), Inches(4.1), Inches(2.0)],
    )

    doc.add_heading("11. Business Metrics and Proof Framework", level=1)
    doc.add_paragraph(
        "APSO must not invent metrics. It should clearly label actual, estimated, placeholder, and AI-inferred values. For the MVP, missing customer proof data should be shown as placeholders until measured."
    )
    table(
        doc,
        ["Metric", "Meaning", "Evidence Source"],
        [
            ["Hours saved", "Estimated or actual effort avoided through earlier detection.", "Finding lifecycle, stakeholder input, historical rework records."],
            ["Rework avoided", "Work prevented by catching gaps before development or release.", "Story gap findings, defect history, resolved actions."],
            ["Defects prevented", "Risks identified before production or QA escalation.", "Quality findings, scanner output, defect records."],
            ["Lead-time reduction", "Improvement in time from ready to done or PR to merge.", "Jira status history, PR timeline, pipeline timeline."],
            ["Coverage improvement", "Increase in test/traceability evidence.", "Coverage reports, test results, PR links."],
            ["Release readiness", "Composite view of risk, tests, pipelines, and unresolved findings.", "Jira, Bitbucket, CI/CD, APSO findings."],
        ],
        [Inches(1.8), Inches(3.25), Inches(3.05)],
    )

    doc.add_heading("12. Production-Grade Business Capabilities", level=1)
    bullets(
        doc,
        [
            "Multi-tenant setup for multiple customers or business units.",
            "Role-based access for executives, managers, product owners, developers, QA, and auditors.",
            "Read-only integration posture for safe onboarding.",
            "Evidence-backed recommendations with source links.",
            "Audit logs for enterprise trust.",
            "Configurable scoring policies in later phases.",
            "Self-hosted/freeware-first deployment posture.",
            "Future optional paid model adapters without changing application behavior.",
        ],
    )

    doc.add_heading("13. Demo Storyline", level=1)
    numbered(
        doc,
        [
            "Start with a Jira story that appears ready.",
            "Show APSO detecting missing acceptance criteria and ambiguous scope.",
            "Open the linked Bitbucket PR/code evidence.",
            "Show quality/security findings generated from deterministic scanner evidence.",
            "Show coverage traceability from Jira story to commits, PRs, tests, and pipeline status.",
            "Show release readiness summary highlighting unresolved risks.",
            "Generate a proof pack explaining what APSO found, what action is needed, and what metric should be captured.",
        ],
    )

    doc.add_heading("14. Stakeholder Acceptance Criteria", level=1)
    bullets(
        doc,
        [
            "Business stakeholders can understand what APSO does without reading code.",
            "Every dashboard and report connects to business outcomes.",
            "AI findings are explainable and evidence-backed.",
            "The product avoids unsafe write-back during MVP.",
            "The MVP can be demonstrated with real read-only Jira/Bitbucket data.",
            "Proof metrics are clearly separated from placeholders.",
            "The product narrative supports a future sellable application.",
        ],
    )

    footer(doc, "APSO Business Functionality and Value Guide | Stakeholder-ready functionality guide")
    doc.save(BUSINESS_DOC)


if __name__ == "__main__":
    technical_doc()
    business_doc()
